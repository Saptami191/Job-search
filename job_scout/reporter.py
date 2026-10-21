import csv
from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .config import load_config, OUTPUT_DIR
from .db import connect
from .guardrails import evaluate_job

def generate_reports(
    limit: int = 30,
    include_unknown: bool = False,
    only_passed_guardrail: bool = True
) -> Dict[str, Path]:
    config = load_config()
    min_score = int(config.get("minimum_score", 15))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with connect() as con:
        rows = [dict(r) for r in con.execute("""
            SELECT * FROM jobs
            ORDER BY COALESCE(NULLIF(ai_score, 0), score) DESC, posted_date DESC
        """).fetchall()]

    shortlist = []
    for row in rows:
        # Check guardrail filter
        if only_passed_guardrail and row.get("guardrail_passed") == 0:
            continue
        effective_score = row.get("ai_score") or row.get("score") or 0
        if effective_score < min_score:
            continue
        
        # Check eligibility terms if not including unknown
        if not include_unknown:
            eligible_raw = str(row.get("eligible_countries") or "").lower()
            allowed_terms = [t.lower() for t in config.get("eligible_country_terms", ["india", "worldwide", "remote"])]
            loc = str(row.get("location") or "").lower()
            if not any(t in eligible_raw for t in allowed_terms) and not any(t in loc for t in allowed_terms):
                continue

        shortlist.append(row)

    shortlist = shortlist[:limit]

    # 1. Generate CSV
    csv_path = OUTPUT_DIR / "ranked_jobs.csv"
    csv_columns = [
        "job_id", "title", "company", "location", "score", "ai_score", "ai_verdict",
        "guardrail_passed", "matched_skills", "status", "url", "source", "posted_date"
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=csv_columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(shortlist)

    # 2. Generate Styled Excel Report
    xlsx_path = OUTPUT_DIR / "ranked_jobs.xlsx"
    write_styled_excel(xlsx_path, shortlist)

    return {
        "csv": csv_path,
        "xlsx": xlsx_path,
        "count": len(shortlist)
    }

def write_styled_excel(path: Path, jobs: List[Dict[str, Any]]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ranked Job Shortlist"

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="0000FF", underline="single")
    
    score_high_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # Soft green
    score_med_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Soft yellow
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        "Rank", "Score", "AI Verdict", "Title", "Company", "Location",
        "Status", "Guardrail", "Matched Skills", "Official URL", "Source", "Date"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for i, j in enumerate(jobs, 1):
        score = j.get("ai_score") or j.get("score") or 0
        verdict = j.get("ai_verdict") or ("RECOMMENDED" if score >= 35 else "BORDERLINE")
        guardrail = "PASS" if j.get("guardrail_passed", 1) == 1 else "FAIL"
        
        row_values = [
            i,
            score,
            verdict,
            j.get("title", ""),
            j.get("company", ""),
            j.get("location", ""),
            j.get("status", "saved"),
            guardrail,
            j.get("matched_skills", ""),
            j.get("url", ""),
            j.get("source", ""),
            j.get("posted_date", "")
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        # Style score cell
        score_cell = ws.cell(row=current_row, column=2)
        score_cell.alignment = Alignment(horizontal="center", vertical="center")
        if score >= 50:
            score_cell.fill = score_high_fill
        elif score >= 25:
            score_cell.fill = score_med_fill

        # Style link cell
        url = j.get("url", "")
        url_cell = ws.cell(row=current_row, column=10)
        if url.startswith("http"):
            url_cell.hyperlink = url
            url_cell.value = "Apply Link"
            url_cell.font = link_font
        
        # Apply borders and alignment
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx not in (2, 10):
                c.font = data_font

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    wb.save(path)
